"""
Pipeline CLI: the .cfg file, the architecture workbook, the pass/fail rule, the
match threshold, explicit source roots, and full runs over the VCU corpus.
"""
import os
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, os.path.abspath("src"))

from Application_Logic import Logic_CallTree_Test as ct
from Application_Logic import Logic_Pipeline as lp
from Application_Logic.Logic_Source_Store import (
    FilesystemSourceProvider, iter_source_roots, root_prefixes,
)
from backend.events import EventBus
from backend.state import AppState
from cli.main import EXIT_ERROR, EXIT_FAILED, EXIT_PASSED, main

VCU = Path("Resources/VCU_TestProject/releases/R5.0").resolve()
REACHED = "Bcm_Lighting_Init"
NOT_CALLED = "Bcm_Lighting_GetMaxCurrentMa"
COLUMNS = ["Port Name", "Operations"]
MAINS = [["Model", "Main Function"], ["Body", "Vcu_InitAll"]]


def _workbook(path, models, mains=MAINS, main_sheet="Main Functions"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in models.items():
        sheet = wb.create_sheet(name)
        for row in rows:
            sheet.append(row)
    sheet = wb.create_sheet(main_sheet)
    for row in mains:
        sheet.append(row)
    wb.save(path)
    return str(path)


def _cfg(tmp_path, *lines, name="run.cfg", defaults=True):
    (tmp_path / "wb.xlsx").touch()
    (tmp_path / "app.out").touch()
    base = ["[inputs]", "workbook = wb.xlsx", "symbols = app.out",
            "[output]", "folder = out"] if defaults else []
    path = tmp_path / name
    path.write_text("\n".join(base + list(lines)) + "\n", encoding="utf-8")
    return str(path)


# ---------------------------------------------------------------------------
# .cfg
# ---------------------------------------------------------------------------
def test_config_resolves_paths_and_lists(tmp_path):
    (tmp_path / "src").mkdir()
    (tmp_path / "gen").mkdir()
    (tmp_path / "fbl.out").touch()
    cfg = lp.load_config(_cfg(tmp_path, "[project]", "release = R1",
                              "[test]", "min_match_score = 85", "max_unmatched = 2"),
                         environ={})
    assert cfg.name == "run" and cfg.release == "R1"
    assert cfg.workbook == str(tmp_path / "wb.xlsx")
    assert cfg.symbols == [("Main", str(tmp_path / "app.out"))]
    assert (cfg.min_match_score, cfg.max_unmatched) == (85, 2)
    assert cfg.password is None


def test_config_multi_line_values_and_environment(tmp_path):
    (tmp_path / "b").mkdir()
    (tmp_path / "b" / "src").mkdir()
    (tmp_path / "b" / "gen").mkdir()
    (tmp_path / "app.out").touch()
    (tmp_path / "fbl.out").touch()
    (tmp_path / "wb.xlsx").touch()
    path = _cfg(tmp_path, "[inputs]", "workbook = wb.xlsx",
                "symbols = APP=app.out", "    FBL=fbl.out",
                "source = ${BUILD}/src", "    ${BUILD}/gen",
                "[output]", "folder = out", defaults=False)
    cfg = lp.load_config(path, environ={"BUILD": str(tmp_path / "b")})
    assert [c for c, _p in cfg.symbols] == ["APP", "FBL"]
    assert cfg.source == [str(tmp_path / "b" / "src"), str(tmp_path / "b" / "gen")]


@pytest.mark.parametrize("lines, message", [
    (["[test]", "min_match_score = 101"], "min_match_score"),
    (["[test]", "max_unmatched = -2"], "max_unmatched"),
    (["[test]", "max_unmatched = two"], "whole number"),
    (["[test]", "strict = maybe"], "strict"),
    (["[test]", "typo = 1"], "Unknown key"),
    (["[typo]", "a = 1"], "Unknown section"),
    (["[project]", "name = ${NOT_SET}"], "NOT_SET"),
    (["[security]", "encryption = maybe"], "encryption"),
])
def test_config_rejects_mistakes(tmp_path, lines, message):
    with pytest.raises(lp.PipelineError, match=message):
        lp.load_config(_cfg(tmp_path, *lines), environ={})


def test_config_reports_missing_paths(tmp_path):
    path = _cfg(tmp_path, "[project]", "name = x")
    os.remove(tmp_path / "app.out")
    with pytest.raises(lp.PipelineError, match="app.out"):
        lp.load_config(path, environ={})


def test_all_three_password_options(tmp_path):
    none = lp.load_config(_cfg(tmp_path, "[security]", "encryption = none"), environ={})
    config = lp.load_config(_cfg(tmp_path, "[security]", "encryption = config",
                                 "password = a;b#c%d"), environ={})
    env = lp.load_config(_cfg(tmp_path, "[security]", "encryption = env"),
                         environ={"ARCH_PASSWORD": "secret"})
    custom = lp.load_config(_cfg(tmp_path, "[security]", "encryption = env",
                                 "password_env = MY_SECRET"), environ={"MY_SECRET": "s2"})
    assert none.password is None
    assert config.password == "a;b#c%d"
    assert env.password == "secret"
    assert custom.password == "s2"


@pytest.mark.parametrize("lines, environ", [
    (["encryption = env"], {}),                         # unset variable
    (["encryption = env"], {"ARCH_PASSWORD": ""}),      # unset masked secret
    (["encryption = config"], {}),
    (["encryption = config", "password = master123"], {}),
    (["encryption = env"], {"ARCH_PASSWORD": "master123"}),
    (["encryption = none", "password = secret"], {}),
])
def test_password_problems_never_fall_back_to_plaintext(tmp_path, lines, environ):
    with pytest.raises(lp.PipelineError):
        lp.load_config(_cfg(tmp_path, "[security]", *lines), environ=environ)


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------
def _config(tmp_path, path, **kw):
    kw.setdefault("function_column", "Operations")
    kw.setdefault("port_column", "Port Name")
    return lp.PipelineConfig(name="x", workbook=path, symbols=[],
                             output_folder=str(tmp_path), **kw)


def test_workbook_sheets_become_models(tmp_path):
    path = _workbook(tmp_path / "w.xlsx", {
        "Body": [COLUMNS + ["Direction"], ["p1", "f1", "in"], [None, None, None]],
        "Power": [COLUMNS, ["p2", "f2"]],
    }, mains=[["M", "F"], ["Body", "main_a"], ["Body", "main_b"], ["Power", "main_c"],
              ["Body", "main_a"]])
    wb = lp.read_workbook(_config(tmp_path, path))
    assert list(wb.models) == ["Body", "Power"]
    assert wb.models["Body"] == [{"Input Port": {"text": "p1"},
                                  "Mapped Func": {"text": "f1"},
                                  "Direction": {"text": "in"}}]
    assert wb.mains == {"Body": ["main_a", "main_b"], "Power": ["main_c"]}


def test_custom_main_sheet_and_extra_column_layout(tmp_path):
    path = _workbook(tmp_path / "w.xlsx", {"Body": [COLUMNS + ["Direction"], ["p", "f", "in"]]},
                     main_sheet="Mains")
    wb = lp.read_workbook(_config(tmp_path, path, main_sheet="Mains"))
    names = [c[0] for c in lp.build_layout(wb.columns)]
    assert names[0] == "TC. ID"
    assert names.index("Direction") == names.index("Review Status") - 1
    assert "Port Name" not in names


@pytest.mark.parametrize("models, mains, message", [
    ({"Body": [["Port Name"], ["p"]]}, MAINS, "no 'Operations' column"),
    ({"Body": [COLUMNS + ["Review Status"], ["p", "f", "x"]]}, MAINS, "reserved"),
    ({"Body": [COLUMNS + ["Mapped Func"], ["p", "f", "x"]]}, MAINS, "same table column"),
    ({"Body": [COLUMNS, ["p", "f"]]}, [["M", "F"], ["Bdoy", "main"]], "no sheet named 'Bdoy'"),
    ({"Body": [COLUMNS, ["p", "f"]]}, [["M", "F"], ["Body", None]], "needs both"),
    ({"Body": [COLUMNS, ["p", "f"]]},
     [["M", "F"]] + [["Body", f"m{i}"] for i in range(ct.MAX_MAINS + 1)], "More than"),
])
def test_workbook_problems_are_setup_errors(tmp_path, models, mains, message):
    path = _workbook(tmp_path / "w.xlsx", models, mains=mains)
    with pytest.raises(lp.PipelineError, match=message):
        lp.read_workbook(_config(tmp_path, path))


def test_workbook_needs_main_sheet_and_models(tmp_path):
    path = _workbook(tmp_path / "w.xlsx", {"Body": [COLUMNS, ["p", "f"]]}, main_sheet="Other")
    with pytest.raises(lp.PipelineError, match="Main Functions"):
        lp.read_workbook(_config(tmp_path, path))
    path = _workbook(tmp_path / "w2.xlsx", {})
    with pytest.raises(lp.PipelineError, match="no model sheets"):
        lp.read_workbook(_config(tmp_path, path))
    with pytest.raises(lp.PipelineError, match=".xlsx"):
        lp.read_workbook(_config(tmp_path, str(tmp_path / "ports.csv")))


# ---------------------------------------------------------------------------
# Pass/fail rule
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("counts, allowed, passed", [
    ({"Passed": 5, "Entry point": 1, "No Result": 2}, 0, True),
    ({"Passed": 5, "Not Matched": 1}, 0, False),
    ({"Passed": 5, "Not Matched": 2}, 2, True),
    ({"Passed": 5, "Not Matched": 3}, 2, False),
    ({"Passed": 5, "Failed": 1}, 99, False),
    ({"Passed": 5, "No Data": 1}, 99, False),
    ({"Passed": 5, "Not Matched": 500}, -1, True),
    ({"Passed": 5, "Failed": 1, "Not Matched": 1}, -1, False),
    ({"Passed": 5, "No Data": 1}, -1, False),
    ({"Not Matched": 5}, -1, False),
    ({"Not Matched": 5}, 5, False),
    ({"No Result": 3}, 0, False),
    ({}, 0, False),
])
def test_gate(counts, allowed, passed):
    gate = lp.evaluate_gate({"M": counts}, allowed)
    assert gate["passed"] is passed
    assert bool(gate["reasons"]) is not passed


def test_gate_adds_up_every_model():
    gate = lp.evaluate_gate({"A": {"Not Matched": 1}, "B": {"Not Matched": 2, "Passed": 1}}, 2)
    assert not gate["passed"]
    assert gate["totals"] == {"Not Matched": 3, "Passed": 1}


# ---------------------------------------------------------------------------
# Match threshold
# ---------------------------------------------------------------------------
CODE_MAP = {"functions": {"main": {"calls": ["motor_stop"]}, "motor_stop": {"calls": []}}}
BINDING = {"release_id": 1, "match_col": "F (Match)", "strict": False}


def _verdict(match_cell, minimum):
    rows = [{"F (Match)": match_cell}]
    ct.run_on_rows(rows, "T", BINDING, CODE_MAP, None, port_state_col=None,
                   tested_at="t", mains=["main"], min_match_score=minimum)
    return rows[0]["T"]


def test_score_at_or_above_threshold_is_a_match():
    assert _verdict({"widget_text": "motor_stop (90%)"}, 90)["text"] == ct.PASSED


def test_score_below_threshold_is_no_match():
    cell = _verdict({"widget_text": "motor_stop (89%)"}, 90)
    assert cell["text"] == ct.NOT_MATCHED
    assert "motor_stop" in cell["reason"] and "89%" in cell["reason"]


def test_name_outside_the_call_tree_is_not_matched():
    cell = _verdict({"widget_text": "some_macro (100%)"}, 90)
    assert cell["text"] == ct.NOT_MATCHED and "call tree" in cell["reason"]


def test_no_threshold_keeps_the_app_behaviour():
    assert _verdict({"widget_text": "motor_stop (20%)"}, None)["text"] == ct.PASSED


def test_hand_picked_match_skips_threshold():
    cell = {"widget_text": "motor_stop (20%)", "user_changed": True}
    assert _verdict(cell, 90)["text"] == ct.PASSED


# ---------------------------------------------------------------------------
# Explicit source roots
# ---------------------------------------------------------------------------
def _tree(root, *files):
    for rel in files:
        path = root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("int x;\n", encoding="utf-8")


def test_listed_folder_overrides_every_skip(tmp_path):
    root = tmp_path / "build"
    _tree(root, "a.c", "output/gen.c", "tools/t.h")
    assert FilesystemSourceProvider(str(root)).list_files() == []   # app default unchanged
    counts = {}
    names = sorted(n for n, _t in iter_source_roots([str(root)], counts=counts))
    assert names == ["a.c", "output/gen.c", "tools/t.h"]
    assert counts == {str(root): 3}


def test_exclude_leaves_out_folders_and_files(tmp_path):
    _tree(tmp_path / "src", "a.c", "stubs/s.c", "b.c")
    names = sorted(n for n, _t in iter_source_roots(
        [str(tmp_path / "src")],
        exclude=[str(tmp_path / "src" / "stubs"), str(tmp_path / "src" / "b.c")]))
    assert names == ["a.c"]


def test_two_roots_with_the_same_path_are_both_kept(tmp_path):
    _tree(tmp_path / "x" / "src", "a.c")
    _tree(tmp_path / "y" / "src", "a.c")
    _tree(tmp_path / "gen", "a.c")
    roots = [str(tmp_path / "x" / "src"), str(tmp_path / "y" / "src"), str(tmp_path / "gen")]
    assert root_prefixes(roots) == ["src", "src_2", "gen"]
    assert sorted(n for n, _t in iter_source_roots(roots)) == \
        ["gen/a.c", "src/a.c", "src_2/a.c"]


# ---------------------------------------------------------------------------
# Full runs over the VCU corpus
# ---------------------------------------------------------------------------
def _run(tmp_path, capsys, monkeypatch, rows, *security, test=("min_match_score = 90",),
         mains=None, env=None):
    monkeypatch.setenv("ARCH_PREFS_FILE", str(tmp_path / "prefs.json"))
    for key, value in (env or {}).items():
        monkeypatch.setenv(key, value)
    mains = mains or [["M", "F"], ["Body", "Vcu_InitAll"]]
    _workbook(tmp_path / "wb.xlsx", {"Body": [COLUMNS] + rows}, mains=mains)
    path = tmp_path / "run.cfg"
    path.write_text("\n".join([
        "[project]", "name = demo", "release = R5",
        "[inputs]", "workbook = wb.xlsx", f"symbols = APP={VCU / 'vcu_R5.0.elf'}",
        f"source = {VCU / 'src'}",
        "[columns]", "function_column = Operations", "port_column = Port Name",
        "[test]", *test,
        "[output]", "folder = out",
        *(("[security]", *security) if security else ()),
    ]) + "\n", encoding="utf-8")
    code = main(["run", "--config", str(path)])
    return code, capsys.readouterr().out


def _open(tmp_path, password=None):
    state = AppState(EventBus())
    state.open_project(str(tmp_path / "out" / "demo.arch"), password=password)
    return state


def test_passing_run(tmp_path, capsys, monkeypatch):
    code, out = _run(tmp_path, capsys, monkeypatch, [["p1", REACHED]])
    assert code == EXIT_PASSED and "RESULT: PASSED" in out
    html = (tmp_path / "out" / "demo.html").read_text(encoding="utf-8")
    assert "Body" in html and REACHED in html
    # The output project opens in the app with the one workbook model, no placeholder.
    state = _open(tmp_path)
    try:
        assert [m.name for m in state.arch_manager.models] == ["Body"]
        assert state.arch_manager.get_active_model().name == "Body"
        assert state.release_manager.get_active_release().name == "R5"
        assert state.db.get_release_symbol_files(1)[0]["component"] == "APP"
        assert state.db.has_release_source(1)
    finally:
        state.close_project()


def test_second_run_replaces_the_output(tmp_path, capsys, monkeypatch):
    assert _run(tmp_path, capsys, monkeypatch, [["p1", REACHED]])[0] == EXIT_PASSED
    assert _run(tmp_path, capsys, monkeypatch, [["p1", REACHED]])[0] == EXIT_PASSED


def test_uncalled_function_fails_and_still_writes_the_report(tmp_path, capsys, monkeypatch):
    code, out = _run(tmp_path, capsys, monkeypatch, [["p1", REACHED], ["p2", NOT_CALLED]])
    assert code == EXIT_FAILED and "RESULT: FAILED" in out
    assert (tmp_path / "out" / "demo.html").exists()
    assert (tmp_path / "out" / "demo.arch").exists()


def test_low_score_is_no_match_and_tolerance_decides(tmp_path, capsys, monkeypatch):
    rows = [["p1", REACHED], ["p2", "Zzqq_Nothing_Like_This"]]
    code, out = _run(tmp_path, capsys, monkeypatch, rows,
                     test=("min_match_score = 95",))
    assert code == EXIT_FAILED and "1 Not Matched" in out
    code, _out = _run(tmp_path, capsys, monkeypatch, rows,
                      test=("min_match_score = 95", "max_unmatched = 1"))
    assert code == EXIT_PASSED


def test_bypass_accepts_any_number_of_unmatched_ports(tmp_path, capsys, monkeypatch):
    rows = [["p1", REACHED], ["p2", "Zzqq_Nothing_Like_This"], ["p3", "Qqzz_Nor_This_One"]]
    code, out = _run(tmp_path, capsys, monkeypatch, rows,
                     test=("min_match_score = 95", "max_unmatched = -1"))
    assert code == EXIT_PASSED and "2 Not Matched" in out


def test_model_without_main_is_no_data_but_the_run_finishes(tmp_path, capsys, monkeypatch):
    code, out = _run(tmp_path, capsys, monkeypatch, [["p1", REACHED]],
                     mains=[["M", "F"]])
    assert code == EXIT_FAILED
    assert "Body has no main function" in out and "1 No Data" in out


def test_encrypted_output_needs_the_password(tmp_path, capsys, monkeypatch):
    from Application_Logic.Logic_Crypto import PasswordRequired
    code, out = _run(tmp_path, capsys, monkeypatch, [["p1", REACHED]],
                     "encryption = env", env={"ARCH_PASSWORD": "pipeline-secret"})
    assert code == EXIT_PASSED
    assert "pipeline-secret" not in out
    with pytest.raises(PasswordRequired):
        AppState(EventBus()).open_project(str(tmp_path / "out" / "demo.arch"))
    state = _open(tmp_path, password="pipeline-secret")
    try:
        assert [m.name for m in state.arch_manager.models] == ["Body"]
    finally:
        state.close_project()


def test_setup_errors_exit_2(tmp_path, capsys, monkeypatch):
    code, out = _run(tmp_path, capsys, monkeypatch, [["p1", REACHED]],
                     "encryption = env", env={"ARCH_PASSWORD": ""})
    assert code == EXIT_ERROR and "ERROR:" in out
    assert not (tmp_path / "out" / "demo.arch").exists()
    assert main(["run", "--config", str(tmp_path / "missing.cfg")]) == EXIT_ERROR


def test_listed_folder_without_source_exits_2(tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("ARCH_PREFS_FILE", str(tmp_path / "prefs.json"))
    _workbook(tmp_path / "wb.xlsx", {"Body": [COLUMNS, ["p1", REACHED]]})
    (tmp_path / "empty").mkdir()
    path = tmp_path / "run.cfg"
    path.write_text("\n".join([
        "[inputs]", "workbook = wb.xlsx", f"symbols = {VCU / 'vcu_R5.0.elf'}",
        "source = empty", "[columns]", "function_column = Operations",
        "port_column = Port Name", "[output]", "folder = out",
    ]) + "\n", encoding="utf-8")
    assert main(["run", "--config", str(path)]) == EXIT_ERROR
    assert "No source files found" in capsys.readouterr().out
